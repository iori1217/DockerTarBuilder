#!/usr/bin/env python3
"""Generate PPT for 用电信息采集系统."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Theme colors - power grid / smart grid style
PRIMARY = RGBColor(0x00, 0x5B, 0x96)      # deep blue
SECONDARY = RGBColor(0x00, 0x8C, 0xBA)    # cyan blue
ACCENT = RGBColor(0xF5, 0xA6, 0x23)       # amber accent
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x2B, 0x3C)
LIGHT_BG = RGBColor(0xE8, 0xF4, 0xFC)
GRAY = RGBColor(0x5A, 0x6A, 0x7A)


def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header_bar(slide, title_text, prs):
    """Add top accent bar and slide title."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(8.8), Inches(0.7))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.font.name = "Microsoft YaHei"


def add_bullet_slide(prs, title, bullets, subtitle=None):
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, title, prs)

    y_start = 1.2 if not subtitle else 1.5
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.05), Inches(8.8), Inches(0.45))
        sp = sub_box.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(14)
        sp.font.color.rgb = GRAY
        sp.font.name = "Microsoft YaHei"

    body = slide.shapes.add_textbox(Inches(0.6), Inches(y_start), Inches(8.8), Inches(5.5))
    tf = body.text_frame
    tf.word_wrap = True

    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = DARK
        p.font.name = "Microsoft YaHei"
        p.space_after = Pt(14)
        p.line_spacing = 1.3

    # left accent line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(y_start), Inches(0.06), Inches(4.5)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = SECONDARY
    line.line.fill.background()

    return slide


def add_title_slide(prs):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide, PRIMARY)

    # decorative shapes
    deco = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(4.8), prs.slide_width, Inches(0.08)
    )
    deco.fill.solid()
    deco.fill.fore_color.rgb = ACCENT
    deco.line.fill.background()

    deco2 = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(7.5), Inches(0), Inches(2.5), prs.slide_height
    )
    deco2.fill.solid()
    deco2.fill.fore_color.rgb = RGBColor(0x00, 0x4A, 0x7C)
    deco2.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6.5), Inches(1.2))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = "用电信息采集系统"
    tp.font.size = Pt(44)
    tp.font.bold = True
    tp.font.color.rgb = WHITE
    tp.font.name = "Microsoft YaHei"

    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(6.5), Inches(0.8))
    sp = sub_box.text_frame.paragraphs[0]
    sp.text = "智能电网重要组成部分 · 业务应用核心数据支撑平台"
    sp.font.size = Pt(20)
    sp.font.color.rgb = RGBColor(0xB8, 0xD4, 0xE8)
    sp.font.name = "Microsoft YaHei"

    date_box = slide.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(4), Inches(0.4))
    dp = date_box.text_frame.paragraphs[0]
    dp.text = "2026"
    dp.font.size = Pt(14)
    dp.font.color.rgb = RGBColor(0x90, 0xB8, 0xD0)
    dp.font.name = "Microsoft YaHei"


def add_overview_slide(prs):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "系统概述", prs)

    # highlight box
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.2), Inches(8.8), Inches(1.6)
    )
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_BG
    box.line.color.rgb = SECONDARY
    box.line.width = Pt(1)

    tb = slide.shapes.add_textbox(Inches(0.9), Inches(1.45), Inches(8.2), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "采集系统全称是用电信息采集系统，是智能电网的重要组成部分，"
        "是各类业务应用的重要数据支撑平台。"
    )
    p.font.size = Pt(20)
    p.font.color.rgb = DARK
    p.font.name = "Microsoft YaHei"
    p.line_spacing = 1.4

    desc = slide.shapes.add_textbox(Inches(0.6), Inches(3.1), Inches(8.8), Inches(2.5))
    dtf = desc.text_frame
    dtf.word_wrap = True
    dp = dtf.paragraphs[0]
    dp.text = (
        "采集系统通过对电力用户的用电信息进行采集、处理和实时监控，"
        "为电力企业生产经营和客户服务提供全面、准确、及时的数据支撑。"
    )
    dp.font.size = Pt(17)
    dp.font.color.rgb = GRAY
    dp.font.name = "Microsoft YaHei"
    dp.line_spacing = 1.4

    # three icon-like labels
    labels = [
        ("智能电网", "核心组成部分"),
        ("数据支撑", "业务应用平台"),
        ("实时采集", "用电信息处理"),
    ]
    for i, (main, sub) in enumerate(labels):
        x = Inches(0.8 + i * 3.0)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(4.5), Inches(2.6), Inches(1.1)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = PRIMARY
        card.line.fill.background()

        mt = slide.shapes.add_textbox(x + Inches(0.15), Inches(4.65), Inches(2.3), Inches(0.45))
        mp = mt.text_frame.paragraphs[0]
        mp.text = main
        mp.font.size = Pt(16)
        mp.font.bold = True
        mp.font.color.rgb = WHITE
        mp.font.name = "Microsoft YaHei"
        mp.alignment = PP_ALIGN.CENTER

        st = slide.shapes.add_textbox(x + Inches(0.15), Inches(5.05), Inches(2.3), Inches(0.4))
        sp = st.text_frame.paragraphs[0]
        sp.text = sub
        sp.font.size = Pt(12)
        sp.font.color.rgb = RGBColor(0xB8, 0xD4, 0xE8)
        sp.font.name = "Microsoft YaHei"
        sp.alignment = PP_ALIGN.CENTER


def add_functions_slide(prs):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "核心功能", prs)

    functions = [
        ("用电信息自动采集", "实现对电力用户用电数据的自动化采集"),
        ("计量异常监测", "及时发现并预警计量装置异常情况"),
        ("电能质量监测", "监测电压、频率等电能质量指标"),
        ("用电分析和管理", "开展用电数据分析与精细化管理"),
        ("相关信息发布", "向用户及相关系统发布用电信息"),
        ("分布式能源监控", "监控分布式光伏、储能等能源设施"),
        ("智能用电设备交互", "实现与智能用电设备的信息交互"),
    ]

    cols = 2
    for i, (title, desc) in enumerate(functions):
        col = i % cols
        row = i // cols
        x = Inches(0.6 + col * 4.5)
        y = Inches(1.15 + row * 1.35)

        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(4.2), Inches(1.15)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG if i % 2 == 0 else WHITE
        card.line.color.rgb = SECONDARY
        card.line.width = Pt(0.75)

        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.15), y + Inches(0.2), Inches(0.18), Inches(0.18))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT
        dot.line.fill.background()

        tt = slide.shapes.add_textbox(x + Inches(0.45), y + Inches(0.12), Inches(3.6), Inches(0.4))
        tp = tt.text_frame.paragraphs[0]
        tp.text = title
        tp.font.size = Pt(15)
        tp.font.bold = True
        tp.font.color.rgb = PRIMARY
        tp.font.name = "Microsoft YaHei"

        dt = slide.shapes.add_textbox(x + Inches(0.45), y + Inches(0.52), Inches(3.6), Inches(0.55))
        dp = dt.text_frame.paragraphs[0]
        dp.text = desc
        dp.font.size = Pt(12)
        dp.font.color.rgb = GRAY
        dp.font.name = "Microsoft YaHei"


def add_interface_slide(prs):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "系统接口与数据流转", prs)

    sub = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(8.8), Inches(0.5))
    sp = sub.text_frame.paragraphs[0]
    sp.text = "用电信息采集系统作为基础数据来源系统，与之有接口的系统繁多"
    sp.font.size = Pt(16)
    sp.font.color.rgb = GRAY
    sp.font.name = "Microsoft YaHei"

    # center hub
    hub = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.5), Inches(2.8), Inches(2.8), Inches(1.0)
    )
    hub.fill.solid()
    hub.fill.fore_color.rgb = PRIMARY
    hub.line.fill.background()

    ht = slide.shapes.add_textbox(Inches(3.6), Inches(3.05), Inches(2.6), Inches(0.6))
    hp = ht.text_frame.paragraphs[0]
    hp.text = "用电信息采集系统"
    hp.font.size = Pt(14)
    hp.font.bold = True
    hp.font.color.rgb = WHITE
    hp.font.name = "Microsoft YaHei"
    hp.alignment = PP_ALIGN.CENTER

    systems = [
        ("营销系统", Inches(0.8), Inches(2.0)),
        ("乡供平台", Inches(7.2), Inches(2.0)),
        ("配网故障抢修", Inches(0.8), Inches(4.5)),
        ("SCADA系统", Inches(7.2), Inches(4.5)),
    ]

    for name, x, y in systems:
        node = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.2), Inches(0.75)
        )
        node.fill.solid()
        node.fill.fore_color.rgb = SECONDARY
        node.line.fill.background()

        nt = slide.shapes.add_textbox(x + Inches(0.1), y + Inches(0.18), Inches(2.0), Inches(0.45))
        np = nt.text_frame.paragraphs[0]
        np.text = name
        np.font.size = Pt(14)
        np.font.bold = True
        np.font.color.rgb = WHITE
        np.font.name = "Microsoft YaHei"
        np.alignment = PP_ALIGN.CENTER

    note = slide.shapes.add_textbox(Inches(0.6), Inches(5.6), Inches(8.8), Inches(0.5))
    np2 = note.text_frame.paragraphs[0]
    np2.text = "为多个业务系统提供统一、可靠的基础用电数据支撑"
    np2.font.size = Pt(14)
    np2.font.color.rgb = GRAY
    np2.font.name = "Microsoft YaHei"
    np2.alignment = PP_ALIGN.CENTER


def add_departments_slide(prs):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "相关业务部门", prs)

    sub = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(8.8), Inches(0.5))
    sp = sub.text_frame.paragraphs[0]
    sp.text = "用电信息采集系统的数据应用广泛，与之相关的业务部门较多"
    sp.font.size = Pt(16)
    sp.font.color.rgb = GRAY
    sp.font.name = "Microsoft YaHei"

    departments = [
        "营销部",
        "营销服务中心（计量中心）",
        "运检部",
        "采集班",
        "装接班",
        "抄表班",
    ]

    for i, dept in enumerate(departments):
        col = i % 3
        row = i // 3
        x = Inches(0.8 + col * 3.0)
        y = Inches(1.8 + row * 1.6)

        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.6), Inches(1.2)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = PRIMARY if i % 2 == 0 else SECONDARY
        card.line.fill.background()

        dt = slide.shapes.add_textbox(x + Inches(0.1), y + Inches(0.35), Inches(2.4), Inches(0.6))
        dp = dt.text_frame.paragraphs[0]
        dp.text = dept
        dp.font.size = Pt(15)
        dp.font.bold = True
        dp.font.color.rgb = WHITE
        dp.font.name = "Microsoft YaHei"
        dp.alignment = PP_ALIGN.CENTER
        dt.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def add_data_scale_slide(prs):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "系统数据规模", prs)

    # big highlight
    highlight = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(8.0), Inches(2.2)
    )
    highlight.fill.solid()
    highlight.fill.fore_color.rgb = PRIMARY
    highlight.line.fill.background()

    main_text = slide.shapes.add_textbox(Inches(1.3), Inches(2.2), Inches(7.4), Inches(1.4))
    mtf = main_text.text_frame
    mtf.word_wrap = True
    mp = mtf.paragraphs[0]
    mp.text = "用电信息采集系统是目前电力公司内\n数据量最大的支撑系统"
    mp.font.size = Pt(28)
    mp.font.bold = True
    mp.font.color.rgb = WHITE
    mp.font.name = "Microsoft YaHei"
    mp.alignment = PP_ALIGN.CENTER
    mp.line_spacing = 1.5

    points = [
        "海量用户用电数据持续采集与存储",
        "支撑多业务系统高频数据调用与分析",
        "对系统性能、稳定性与扩展性提出极高要求",
    ]
    for i, pt in enumerate(points):
        y = Inches(4.3 + i * 0.65)
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), y + Inches(0.08), Inches(0.15), Inches(0.15))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT
        dot.line.fill.background()

        pt_box = slide.shapes.add_textbox(Inches(1.55), y, Inches(7.5), Inches(0.5))
        pp = pt_box.text_frame.paragraphs[0]
        pp.text = pt
        pp.font.size = Pt(16)
        pp.font.color.rgb = DARK
        pp.font.name = "Microsoft YaHei"


def add_summary_slide(prs):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide, PRIMARY)
    add_header_bar(slide, "总结", prs)

    # fix header color on dark bg
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.paragraphs[0].text == "总结":
            shape.text_frame.paragraphs[0].font.color.rgb = WHITE

    items = [
        "智能电网核心基础设施，业务数据中枢",
        "功能全面：采集、监测、分析、交互一体化",
        "广泛对接营销、配网、SCADA 等系统",
        "服务营销、运检等多业务部门协同作业",
        "电力公司数据量最大的支撑系统",
    ]

    for i, item in enumerate(items):
        y = Inches(1.5 + i * 0.85)
        bullet = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), y + Inches(0.1), Inches(0.2), Inches(0.2))
        bullet.fill.solid()
        bullet.fill.fore_color.rgb = ACCENT
        bullet.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(1.4), y, Inches(7.5), Inches(0.6))
        tp = tb.text_frame.paragraphs[0]
        tp.text = item
        tp.font.size = Pt(20)
        tp.font.color.rgb = WHITE
        tp.font.name = "Microsoft YaHei"

    thanks = slide.shapes.add_textbox(Inches(0.6), Inches(5.8), Inches(8.8), Inches(0.5))
    thp = thanks.text_frame.paragraphs[0]
    thp.text = "谢谢"
    thp.font.size = Pt(24)
    thp.font.bold = True
    thp.font.color.rgb = ACCENT
    thp.font.name = "Microsoft YaHei"
    thp.alignment = PP_ALIGN.CENTER


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)
    add_overview_slide(prs)
    add_functions_slide(prs)
    add_interface_slide(prs)
    add_departments_slide(prs)
    add_data_scale_slide(prs)
    add_summary_slide(prs)

    output_path = "/workspace/用电信息采集系统.pptx"
    prs.save(output_path)
    print(f"PPT saved to: {output_path}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
